"""Insert the captured screenshots into docs/Work1_Report.docx.

Each grey "Screenshot: ..." placeholder is replaced by the corresponding image
plus a small caption. Run after build_report_docx.py.
"""

import os
import shutil

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.text.paragraph import Paragraph

BASE = r"D:\hdu-work1"
SHOTS = r"D:\[HDU-CS4-26-27-1] Information Security\work1\screenshots"
DOCX = os.path.join(BASE, "docs", "Work1_Report.docx")
ROOT_COPY = r"D:\[HDU-CS4-26-27-1] Information Security\work1\Work1_Report.docx"

# order of the placeholders in the document
IMAGES = [
    ("01-api-smoke-test.png", "Figure 1 - API smoke test: 401 without a token, 401 on a wrong password, 200 with a valid JWT"),
    ("02-actions-overview.png", "Figure 2 - GitHub Actions run #8: success, all three jobs green (Tests 17s, SAST 11s, SCA 25s)"),
    ("03-sast-bandit.png", "Figure 3 - GitHub Actions SAST job log: bandit reports no issues, 503 lines scanned"),
    ("04-sca-pip-audit.png", "Figure 4 - GitHub Actions SCA job log: pip-audit reports no known vulnerabilities"),
    ("05-cve-vulnerable.png", "Figure 5 - CVE-2014-3120 reproduced on 1.1.1: the script runs, id is root"),
    ("06-default-config.png", "Figure 6 - the default 1.1.1 config has no script.disable_dynamic, so scripting is on"),
    ("07-fixed.png", "Figure 7 - 1.4.4 rejects the script (dynamic scripting for [groovy] disabled)"),
    ("08-hardened.png", "Figure 8 - hardened 1.1.1 rejects the script, the search still works"),
]


def add_caption(after_paragraph, text):
    new_p = OxmlElement("w:p")
    after_paragraph._p.addnext(new_p)
    paragraph = Paragraph(new_p, after_paragraph._parent)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(10)
    run = paragraph.add_run(text)
    run.italic = True
    run.font.size = Pt(8.5)
    run.font.name = "Times New Roman"
    run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)


def main():
    doc = Document(DOCX)
    placeholders = [p for p in doc.paragraphs if p.text.strip().startswith("Screenshot:")]
    if len(placeholders) != len(IMAGES):
        raise SystemExit(f"expected {len(IMAGES)} placeholders, found {len(placeholders)}")

    for paragraph, (filename, caption) in zip(placeholders, IMAGES):
        path = os.path.join(SHOTS, filename)
        if not os.path.exists(path):
            raise SystemExit(f"missing screenshot: {path}")

        for run in list(paragraph.runs):
            run._element.getparent().remove(run._element)

        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(6)
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.add_run().add_picture(path, width=Cm(15.4))
        add_caption(paragraph, caption)

    doc.save(DOCX)
    shutil.copyfile(DOCX, ROOT_COPY)
    print(f"inserted {len(IMAGES)} images -> {DOCX}")
    print(f"copied -> {ROOT_COPY}")


if __name__ == "__main__":
    main()
