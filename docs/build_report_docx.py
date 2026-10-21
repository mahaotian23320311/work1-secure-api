"""Build docs/Work1_Report.docx from docs/Work1_Report.md.

A tiny Markdown subset is supported: headings (#, ##, ###), paragraphs,
bullet and numbered lists, pipe tables, fenced code blocks, **bold** and
`inline code`. The title page is generated separately so the report can be
handed in without further layout work.

Usage:
    python docs/build_report_docx.py
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).resolve().parent
SOURCE = HERE / "Work1_Report.md"
TARGET = HERE / "Work1_Report.docx"

BODY_FONT = "Times New Roman"
MONO_FONT = "Consolas"


# --------------------------------------------------------------------------- #
# helpers
# --------------------------------------------------------------------------- #
def shade(paragraph, hex_fill: str = "F2F2F2") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    p_pr.append(shd)


def add_runs(paragraph, text: str, *, size: float, mono: bool = False) -> None:
    """Render **bold**, *italic* and `code` spans inside one paragraph."""
    pattern = re.compile(r"(\*\*.+?\*\*|`[^`]+`|(?<![*\w])\*[^*\n]+?\*(?![*\w]))")
    for chunk in pattern.split(text):
        if not chunk:
            continue
        if chunk.startswith("**") and chunk.endswith("**") and len(chunk) > 4:
            run = paragraph.add_run(chunk[2:-2])
            run.bold = True
        elif chunk.startswith("`") and chunk.endswith("`") and len(chunk) > 2:
            run = paragraph.add_run(chunk[1:-1])
            run.font.name = MONO_FONT
            run._element.rPr.rFonts.set(qn("w:eastAsia"), MONO_FONT)
            run.font.size = Pt(size - 1)
            continue
        elif chunk.startswith("*") and chunk.endswith("*") and len(chunk) > 2:
            run = paragraph.add_run(chunk[1:-1])
            run.italic = True
        else:
            run = paragraph.add_run(chunk)
        run.font.size = Pt(size)
        if not mono:
            run.font.name = BODY_FONT


def title_page(doc: Document) -> None:
    def line(text: str, size: float, *, bold: bool = False, space_after: int = 6,
             color: RGBColor | None = None) -> None:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(space_after)
        run = p.add_run(text)
        run.bold = bold
        run.font.size = Pt(size)
        run.font.name = BODY_FONT
        if color is not None:
            run.font.color.rgb = color

    line("HANGZHOU DIANZI UNIVERSITY", 12, bold=True, space_after=2)
    line("Information Security Laboratory", 11, space_after=60)

    line("LABORATORY WORK No. 1", 20, bold=True, space_after=10)
    line("Secure REST API Development with CI/CD", 14, bold=True, space_after=2)
    line("and Real CVE Vulnerability Analysis", 14, bold=True, space_after=60)

    table = doc.add_table(rows=5, cols=2)
    table.style = "Table Grid"
    rows = [
        ("Variant No.", "____________"),
        ("Student (Surname, Name)", "____________"),
        ("HDU number (student ID)", "____________"),
        ("Group", "____________"),
        ("Date", "______ / ______ / 2026"),
    ]
    for row, (label, value) in zip(table.rows, rows):
        row.cells[0].text = label
        row.cells[1].text = value
        for index, cell in enumerate(row.cells):
            for paragraph in cell.paragraphs:
                if index == 1:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    run.font.size = Pt(11)
                    run.font.name = BODY_FONT
    table.columns[0].width = Cm(6.5)
    table.columns[1].width = Cm(7.5)

    doc.add_paragraph()
    line("Repository: https://github.com/mahaotian23320311/work1-secure-api", 10, space_after=2)
    line("Pipeline:   https://github.com/mahaotian23320311/work1-secure-api/actions", 10)

    doc.add_page_break()


# --------------------------------------------------------------------------- #
# markdown -> docx
# --------------------------------------------------------------------------- #
def _join_continuations(lines: list[str]) -> list[str]:
    """Fold indented continuation lines into the list item they belong to."""
    out: list[str] = []
    for raw in lines:
        line = raw.rstrip()
        stripped = line.strip()
        if line.startswith(" ") and stripped and not stripped.startswith(("|", "```")):
            if out and re.match(r"^([-*]|\d+\.)\s", out[-1].strip()):
                out[-1] = out[-1].strip() + " " + stripped
                continue
        out.append(line)
    return out


def convert(md: str, doc: Document) -> None:
    lines = _join_continuations(md.splitlines())

    # Skip the Markdown title page; the docx has its own.
    try:
        start = next(i for i, line in enumerate(lines) if line.startswith("## 1."))
    except StopIteration:
        start = 0

    buffer: list[str] = []

    def flush() -> None:
        """Emit the accumulated lines as one wrapped paragraph."""
        if not buffer:
            return
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(6)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        add_runs(paragraph, " ".join(buffer), size=11)
        buffer.clear()

    i = start
    while i < len(lines):
        line = lines[i].rstrip()

        # Placeholders may be written as `[[SCREENSHOT: ...]]`; unwrap them so
        # they render as a grey note instead of inline code.
        if re.fullmatch(r"`\[\[SCREENSHOT:.*\]\]`", line.strip()):
            line = line.strip().strip("`")

        if not line.strip():
            flush()
            i += 1
            continue

        if line.startswith("```"):
            flush()
            i += 1
            block: list[str] = []
            while i < len(lines) and not lines[i].startswith("```"):
                block.append(lines[i].rstrip())
                i += 1
            i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.4)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(8)
            shade(p)
            run = p.add_run("\n".join(block))
            run.font.name = MONO_FONT
            run._element.rPr.rFonts.set(qn("w:eastAsia"), MONO_FONT)
            run.font.size = Pt(8.5)
            continue

        if line.startswith("|"):
            flush()
            rows: list[list[str]] = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            table = doc.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = "Table Grid"
            for r, cells in enumerate(rows):
                for c, text in enumerate(cells):
                    cell = table.cell(r, c)
                    cell.text = ""
                    paragraph = cell.paragraphs[0]
                    paragraph.paragraph_format.space_after = Pt(2)
                    paragraph.paragraph_format.space_before = Pt(2)
                    add_runs(paragraph, text, size=9)
                    if r == 0:
                        for run in paragraph.runs:
                            run.bold = True
            doc.add_paragraph()
            continue

        if line.startswith("#"):
            flush()
            level = len(line) - len(line.lstrip("#"))
            text = line[level:].strip()
            heading = doc.add_heading(level=min(level, 3) - 1, text=text)
            for run in heading.runs:
                run.font.color.rgb = RGBColor(0, 0, 0)
                run.font.name = BODY_FONT
                run.font.size = Pt({1: 16, 2: 13, 3: 11.5}.get(level, 11.5))
            i += 1
            continue

        if re.match(r"^[-*] ", line):
            flush()
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, line[2:].strip(), size=11)
            p.paragraph_format.space_after = Pt(2)
            i += 1
            continue

        if re.match(r"^\d+\. ", line):
            flush()
            match = re.match(r"^(\d+)\.\s+(.*)$", line.strip())
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.75)
            p.paragraph_format.first_line_indent = Cm(-0.75)
            p.paragraph_format.space_after = Pt(2)
            add_runs(p, match.group(1) + ".\t" + match.group(2), size=11)
            i += 1
            continue

        if line.startswith("[[SCREENSHOT:"):
            flush()
            p = doc.add_paragraph()
            run = p.add_run(line.strip("[]").replace("SCREENSHOT:", "Screenshot:"))
            run.italic = True
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
            i += 1
            continue

        if line.strip() == "---":
            flush()
            i += 1
            continue

        buffer.append(line.strip())
        i += 1

    flush()


def main() -> None:
    doc = Document()

    section = doc.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21.0)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.0)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)

    style = doc.styles["Normal"]
    style.font.name = BODY_FONT
    style.font.size = Pt(11)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")
    style.paragraph_format.line_spacing = 1.0
    style.paragraph_format.space_after = Pt(6)

    title_page(doc)
    convert(SOURCE.read_text(encoding="utf-8"), doc)
    doc.save(TARGET)
    print(f"wrote {TARGET} ({TARGET.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
