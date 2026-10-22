"""Normalise the title page of Work1_Report.docx.

- removes the "Group" row (this course does not use groups),
- fills variant / student / HDU number / date by matching the row label,
  so the script keeps working if the table layout changes.

Safe to run again after build_report_docx.py.
"""

import shutil

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

DOCX = r"D:\hdu-work1\docs\Work1_Report.docx"
ROOT_COPY = r"D:\[HDU-CS4-26-27-1] Information Security\work1\Work1_Report.docx"

VALUES = {
    "Variant": "CVE-2014-3120",
    "Student": "mahaotian",
    "HDU": "23320311",
    "Date": "22 / 10 / 2026",
}

DROP_LABELS = ("Group",)


def main():
    doc = Document(DOCX)
    table = doc.tables[0]

    for row in list(table.rows):
        label = row.cells[0].text.strip()
        if label.startswith(DROP_LABELS):
            row._element.getparent().remove(row._element)

    for row in table.rows:
        label = row.cells[0].text.strip()
        for key, value in VALUES.items():
            if not label.startswith(key):
                continue
            cell = row.cells[1]
            for paragraph in cell.paragraphs:
                for run in list(paragraph.runs):
                    run._element.getparent().remove(run._element)
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run(value)
            run.font.name = "Times New Roman"
            run.font.size = Pt(11)

    doc.save(DOCX)
    shutil.copyfile(DOCX, ROOT_COPY)
    print("rows now:", len(table.rows))
    for row in table.rows:
        print("  ", row.cells[0].text.strip(), "=", row.cells[1].text.strip())
    print("copied ->", ROOT_COPY)


if __name__ == "__main__":
    main()
